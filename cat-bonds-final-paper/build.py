"""Build paper.docx and paper.pdf from paper.md.

Usage: python3 build.py
Needs pandoc, python-docx and LibreOffice (soffice) for the PDF.
"""
import copy
import re
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
SRC = HERE / "paper.md"
DOCX = HERE / "paper.docx"

FONT = "Times New Roman"
BODY_PT = 12
LINE_SPACING = 1.15

COVER = {
    "kicker": "Final paper",
    "title": "Are Catastrophe Bonds Really Zero-Beta Assets?",
    "subtitle": "Diversification Benefits in Normal Markets and Financial Crises",
    "submitted_to": "Submitted to: Mohamed Mnasri",
    "course": "FINA60218A",
    "by": "By:",
    "author": "Tran Nguyen Quoc-Long",
    "date": "December 2026",
}


def set_font(style, size=None, bold=None, italic=None, color=None):
    style.font.name = FONT
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    style.font.color.rgb = color or RGBColor(0, 0, 0)


def style_by_name(doc, name):
    for st in doc.styles:
        if st.name == name:
            return st
    raise KeyError(name)


def make_reference(path):
    default = subprocess.run(
        ["pandoc", "--print-default-data-file", "reference.docx"],
        check=True, capture_output=True,
    ).stdout
    path.write_bytes(default)
    doc = Document(path)
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name in [s.name for s in doc.styles]:
            st = style_by_name(doc, name)
            set_font(st, BODY_PT)
            pf = st.paragraph_format
            pf.line_spacing = LINE_SPACING
            pf.space_before = Pt(0)
            pf.space_after = Pt(4)
            if name in ("Body Text", "First Paragraph"):
                pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name, size in (("Heading 1", 14), ("Heading 2", 12)):
        st = style_by_name(doc, name)
        set_font(st, size, bold=True, italic=False)
        st.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 9)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True
    for name in ("Title", "Author", "Hyperlink"):
        if name in [s.name for s in doc.styles]:
            set_font(style_by_name(doc, name))
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(27.94), Cm(21.59)  # US Letter
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.5))
    doc.save(path)


def add_page_number_footer(section):
    section.different_first_page_header_footer = True
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    run.font.name = FONT
    run.font.size = Pt(11)


def add_cover(doc):
    body = doc.element.body
    first = body[0]
    lines = [
        (COVER["kicker"], 13, False, 60),
        (COVER["title"], 20, True, 18),
        (COVER["subtitle"], 15, False, 90),
        (COVER["submitted_to"], 13, False, 0),
        (COVER["course"], 13, False, 60),
        (COVER["by"], 13, False, 0),
        (COVER["author"], 13, True, 60),
        (COVER["date"], 13, False, 0),
    ]
    new = []
    for text, size, bold, after in lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(after)
        r = p.add_run(text)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.bold = bold
        new.append(p)
    new[0].paragraph_format.space_before = Pt(120)
    brk = doc.add_paragraph()
    brk.add_run().add_break(WD_BREAK.PAGE)
    new.append(brk)
    for p in new:
        first.addprevious(p._p)


def format_tables(doc):
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl_pr = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), "808080")
            borders.append(el)
        old = tbl_pr.find(qn("w:tblBorders"))
        if old is not None:
            tbl_pr.remove(old)
        tbl_pr.append(borders)
        width = OxmlElement("w:tblW")
        width.set(qn("w:w"), "5000")
        width.set(qn("w:type"), "pct")
        old = tbl_pr.find(qn("w:tblW"))
        if old is not None:
            tbl_pr.remove(old)
        tbl_pr.append(width)
        nxt = table._tbl.getnext()
        if nxt is not None and nxt.tag == qn("w:p"):
            ppr = nxt.get_or_add_pPr()
            spacing = OxmlElement("w:spacing")
            spacing.set(qn("w:before"), "160")
            ppr.append(spacing)
        ncols = len(table.columns)
        first_w = 0.16 if ncols == 4 else 0.24
        rest = (1 - first_w) / (ncols - 1)
        total = Cm(21.59 - 5.0)
        for row_i, row in enumerate(table.rows):
            tr_pr = row._tr.get_or_add_trPr()
            tr_pr.append(OxmlElement("w:cantSplit"))
            if row_i == 0:
                tr_pr.append(OxmlElement("w:tblHeader"))
            for col_i, cell in enumerate(row.cells):
                cell.width = int(total * (first_w if col_i == 0 else rest))
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.0
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for r in p.runs:
                        r.font.name = FONT
                        r.font.size = Pt(9.5)
                        if row_i == 0 or col_i == 0:
                            r.bold = True


def format_equations(doc):
    text_width = 21.59 - 5.0
    for p in doc.paragraphs:
        if p.style.name != "Equation" or not p.runs:
            continue
        pf = p.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.space_before = Pt(6)
        pf.space_after = Pt(6)
        pf.tab_stops.add_tab_stop(Cm(text_width / 2), WD_TAB_ALIGNMENT.CENTER)
        pf.tab_stops.add_tab_stop(Cm(text_width), WD_TAB_ALIGNMENT.RIGHT)
        last = p.runs[-1]
        m = re.search(r"\s*(\(\d\))\s*$", last.text)
        if m:
            last.text = last.text[: m.start()] + "\t" + m.group(1)
        first = p.runs[0]
        first.text = "\t" + first.text
        for r in p.runs:
            r.font.name = FONT
            r.font.size = Pt(BODY_PT)


def style_captions_and_refs(doc):
    in_refs = False
    for p in doc.paragraphs:
        text = p.text.strip()
        if p.style.name.startswith("Heading") and text.lower().startswith("references"):
            in_refs = True
            continue
        if text.startswith("Table ") and p.runs and p.runs[0].bold:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(6)
        if in_refs and text:
            pf = p.paragraph_format
            pf.left_indent = Cm(1.0)
            pf.first_line_indent = Cm(-1.0)
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf.space_after = Pt(4)
            pf.line_spacing = 1.0
            for r in p.runs:
                r.font.size = Pt(10.5)


def main():
    text = SRC.read_text(encoding="utf-8")
    text = re.sub(r"\A---\n.*?\n---\n", "", text, flags=re.S)
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        ref = tmp / "reference.docx"
        make_reference(ref)
        md = tmp / "body.md"
        md.write_text(text, encoding="utf-8")
        subprocess.run(
            ["pandoc", str(md), "-f", "markdown", "-t", "docx",
             "--reference-doc", str(ref), "-o", str(DOCX)],
            check=True,
        )
    doc = Document(DOCX)
    add_cover(doc)
    format_tables(doc)
    format_equations(doc)
    style_captions_and_refs(doc)
    add_page_number_footer(doc.sections[0])
    doc.save(DOCX)
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(HERE), str(DOCX)],
        check=True, capture_output=True,
    )
    print("built", DOCX.name, "and paper.pdf")


if __name__ == "__main__":
    main()
